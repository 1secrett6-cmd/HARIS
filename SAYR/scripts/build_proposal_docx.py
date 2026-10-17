"""Build proposal/SAYR_Proposal.docx from docs/PROPOSAL_AR.md (Arabic, right-to-left)."""

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "docs" / "PROPOSAL_AR.md"
OUT = ROOT / "proposal" / "SAYR_Proposal.docx"

FONT = "Arial"
BRAND = RGBColor(0x0B, 0x6E, 0x4F)
ARABIC = re.compile(r"[؀-ۿ]")
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def is_rtl(text):
    return bool(ARABIC.search(text))


def set_bidi(paragraph, rtl):
    ppr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1" if rtl else "0")
    ppr.append(bidi)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT


def style_run(run, rtl, size=None, bold=None, italic=None, color=None, mono=False):
    font = "Consolas" if mono else FONT
    run.font.name = font
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), font)
    if rtl:
        rpr.append(OxmlElement("w:rtl"))
    if size:
        run.font.size = Pt(size)
        szcs = OxmlElement("w:szCs")
        szcs.set(qn("w:val"), str(int(size * 2)))
        rpr.append(szcs)
    if bold:
        run.bold = True
        rpr.append(OxmlElement("w:bCs"))
    if italic:
        run.italic = True
    if color:
        run.font.color.rgb = color


def add_inline(paragraph, text, rtl, size=11, color=None, bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        b, i, mono = bold, False, False
        if part.startswith("**") and part.endswith("**"):
            part, b = part[2:-2], True
        elif part.startswith("`") and part.endswith("`"):
            part, mono = part[1:-1], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, i = part[1:-1], True
        style_run(paragraph.add_run(part), rtl, size=size, bold=b, italic=i, color=color, mono=mono)


def paragraph(doc, text, size=11, color=None, bold=False, style=None, force_rtl=None):
    rtl = is_rtl(text) if force_rtl is None else force_rtl
    p = doc.add_paragraph(style=style)
    set_bidi(p, rtl)
    add_inline(p, text, rtl, size=size, color=color, bold=bold)
    return p


def shade(cell, hex_color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcpr.append(shd)


def add_table(doc, rows):
    rtl = any(is_rtl(c) for r in rows for c in r)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if rtl:
        tblpr = table._tbl.tblPr
        bidi = OxmlElement("w:bidiVisual")
        tblpr.append(bidi)
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            p = cell.paragraphs[0]
            cell_rtl = is_rtl(text) or (rtl and not text.strip())
            set_bidi(p, cell_rtl)
            header = r == 0 and any(x.strip() for x in rows[0])
            add_inline(p, text, cell_rtl, size=10, bold=header,
                       color=RGBColor(0xFF, 0xFF, 0xFF) if header else None)
            if header:
                shade(cell, "0B6E4F")
    doc.add_paragraph()


def parse_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def build():
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(60)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                p = doc.add_paragraph()
                set_bidi(p, False)
                style_run(p.add_run(lines[i]), False, size=8, mono=True)
                p.paragraph_format.space_after = Pt(0)
                i += 1
        elif stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.fullmatch(r"\|[\s:|-]+\|", lines[i].strip()):
                    rows.append(parse_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue
        elif stripped.startswith("# "):
            paragraph(doc, stripped[2:], size=30, color=BRAND, bold=True)
        elif stripped.startswith("## "):
            paragraph(doc, stripped[3:], size=16, color=BRAND, bold=True)
        elif stripped.startswith("### "):
            paragraph(doc, stripped[4:], size=13, color=BRAND, bold=True)
        elif stripped == "---":
            pass
        elif stripped.startswith("> "):
            paragraph(doc, stripped[2:], size=11, color=RGBColor(0x44, 0x44, 0x44))
        elif re.match(r"^[-*] ", stripped):
            paragraph(doc, "•  " + stripped[2:])
        elif re.match(r"^\d+\. ", stripped):
            paragraph(doc, stripped)
        elif stripped:
            paragraph(doc, stripped)
        i += 1

    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
